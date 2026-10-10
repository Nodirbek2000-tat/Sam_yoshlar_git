"""Samarqand startuplar ofisi reestrini Excel'dan yuklash.

Faqat kerakli ustunlar o'qiladi — qolganlari (ijtimoiy holati, so'ragan
mablag'i, batafsil, natija va h.k.) e'tiborsiz qoldiriladi:

    Hudud, F.I.Sh, Yoshi (yoki tug'ilgan sana), Rasmi, Loyiha haqida,
    StartUp sohasi, StartUp bosqichi, Loyiha rasmi, Telefon, Telegram (bo'lsa)

Ustunlar sarlavhasi bo'yicha topiladi — tartibi o'zgarsa ham ishlaydi.
Rasmlar Excel ichiga joylangan bo'ladi: qaysi katakka qo'yilgan bo'lsa,
o'sha qatorniki hisoblanadi.

Bir odam bir nechta loyiha bilan kelsa — har biri alohida yoziladi.
Aynan o'sha odamning aynan o'sha loyihasi qayta yuklansa — o'tkazib yuboriladi.
"""

import hashlib
import re
from datetime import date, datetime

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.core.constants import SamarqandDistrict

from .models import OfficeSphere, OfficeStage, OfficeStartup

#: Bir martada shuncha qatordan ko'p bo'lsa — fayl xato deb hisoblanadi
MAX_ROWS = 3000


def _norm(text):
    text = str(text or '').strip().lower()
    text = re.sub(r"['`’ʻʼ‘´]", '', text)
    return re.sub(r'\s+', ' ', text)


def _clean(text):
    """Bir nechta bo'shliq va qator tashlashlarni bittaga."""
    return re.sub(r'\s+', ' ', str(text or '')).strip()


# -------------------------------------------------------------- ustunlar

def _column_kind(header):
    h = _norm(header)
    if not h:
        return None
    if 'loyiha rasm' in h:
        return 'project_image'
    if 'rasm' in h or 'foto' in h or 'surat' in h:
        return 'photo'
    if 'loyiha nomi' in h or 'startup nomi' in h or 'startap nomi' in h:
        return 'name'
    if 'loyiha haqida' in h or 'tavsif' in h:
        return 'about'
    if 'f.i.sh' in h or 'fish' in h or 'f.i.o' in h or h in ('fio', 'ism', 'ismi', 'asoschi'):
        return 'full_name'
    if 'hudud' in h or 'tuman' in h:
        return 'district'
    if 'yosh' in h or 'tugilgan' in h:
        return 'age'
    if 'soha' in h or 'yonalish' in h:
        return 'sphere'
    if 'bosqich' in h:
        return 'stage'
    if 'telegram' in h:
        return 'telegram'
    if 'telefon' in h:
        return 'phone'
    return None


def _find_header(ws):
    """Sarlavha qatori va {tur: ustun raqami}. F.I.Sh ustuni bo'lishi shart."""
    for row in range(1, min(ws.max_row, 6) + 1):
        columns = {}
        for cell in ws[row]:
            kind = _column_kind(cell.value)
            if kind and kind not in columns:
                columns[kind] = cell.column
        if 'full_name' in columns:
            return row, columns
    return None, {}


# -------------------------------------------------------------- qiymatlar

SPHERE_WORDS = [
    (OfficeSphere.EDTECH, ['edtech', 'talim']),
    (OfficeSphere.GREENTECH, ['greentech', 'ekolog', 'energiya']),
    (OfficeSphere.AI, ['ai & data', 'ai &', 'suniy intellekt', 'ai/data']),
    (OfficeSphere.HEALTHTECH, ['healthtech', 'medtech', 'tibbiyot']),
    (OfficeSphere.SOCIAL, ['social', 'ijtimoiy']),
    (OfficeSphere.FINTECH, ['fintech', 'moliya']),
    (OfficeSphere.AGRITECH, ['agritech', 'agrotech', 'qishloq']),
    (OfficeSphere.ECOMMERCE, ['e-commerce', 'ecommerce', 'savdo']),
    (OfficeSphere.LOGISTICS, ['logistic', 'logistika', 'transport']),
    (OfficeSphere.SMARTCITY, ['smart city', 'aqlli shahar']),
    (OfficeSphere.TOURISM, ['turizm', 'tourism', 'sayohat']),
]


def parse_sphere(raw):
    text = _norm(raw)
    for value, words in SPHERE_WORDS:
        if any(word in text for word in words):
            return value
    return OfficeSphere.OTHER


def parse_stage(raw):
    text = _norm(raw)
    if not text:
        return OfficeStage.IDEA
    if 'mvp' in text:
        return OfficeStage.MVP
    if 'seed' in text:
        return OfficeStage.SEED
    if 'sotuv' in text or 'biznes' in text or 'kengay' in text:
        return OfficeStage.SALES
    return OfficeStage.IDEA


def parse_district(raw):
    """«Samarqand shahar», «Urgut tumani», «urgut» -> kod. Topilmasa — matnning o'zi."""
    text = _norm(raw).replace(' shahar', ' shahri').replace(' sh.', ' shahri')
    if not text:
        return ''
    for value, label in SamarqandDistrict.choices:
        if text in (value, _norm(label)):
            return value
    shorts = {}
    for value, label in SamarqandDistrict.choices:
        short = _norm(label).replace(' tumani', '').replace(' shahri', '')
        shorts.setdefault(short, []).append(value)
    found = shorts.get(text)
    if found and len(found) == 1:
        return found[0]
    return _clean(raw)[:40]


def parse_age(raw):
    """(tug'ilgan sana, yosh). Excel'da sana, «13.06.2000» yoki «22» bo'lishi mumkin."""
    if raw is None or raw == '':
        return None, None
    if isinstance(raw, datetime):
        return raw.date(), None
    if isinstance(raw, date):
        return raw, None
    if isinstance(raw, (int, float)) and 5 <= raw <= 99:
        return None, int(raw)
    text = _clean(raw)
    match = re.fullmatch(r'(\d{1,2})[./-](\d{1,2})[./-](\d{4})', text)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year, month, day), None
        except ValueError:
            return None, None
    match = re.fullmatch(r'(\d{4})[./-](\d{1,2})[./-](\d{1,2})', text)
    if match:
        year, month, day = map(int, match.groups())
        try:
            return date(year, month, day), None
        except ValueError:
            return None, None
    if text.isdigit() and 5 <= int(text) <= 99:
        return None, int(text)
    return None, None


PHONE = re.compile(r'(?:\+?998[\s()-]*)?\d{2}[\s()-]*\d{3}[\s-]*\d{2}[\s-]*\d{2}')


def parse_phone(raw):
    """Birinchi telefon raqami (katakda ikkita yozilgan bo'lishi mumkin) -> «+998901234567»."""
    match = PHONE.search(str(raw or ''))
    digits = re.sub(r'\D', '', match.group(0) if match else str(raw or ''))
    if len(digits) == 9:
        digits = '998' + digits
    if len(digits) == 12 and digits.startswith('998'):
        return '+' + digits
    return _clean(raw)[:25]


def parse_telegram(raw):
    value = _clean(raw)
    value = re.sub(r'^(https?://)?(t\.me|telegram\.me)/', '', value, flags=re.IGNORECASE)
    value = value.lstrip('@').split('?')[0].strip('/')
    return value if re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{3,31}', value) else ''


SEPARATORS = re.compile(r'\s+[—–-]\s+|[—–]|\s*-\s+bu\b|-\s')


def project_name(about):
    """Loyiha matnining boshidan nomi: «Samly- bu CRM…» -> «Samly»."""
    text = _clean(about)
    if not text:
        return ''
    match = SEPARATORS.search(text)
    if match and 2 <= match.start() <= 48:
        head = text[:match.start()].strip(' :.,')
        # «Muammo: …», «Hozirda …» kabi jumla boshlari nom emas
        if head and not head.lower().startswith(('muammo', 'hozirda', 'loyiha')):
            return head
    comma = text.find(',')
    if 2 <= comma <= 40 and not text.lower().startswith(('muammo', 'hozirda')):
        return text[:comma].strip()
    if len(text) <= 60:
        return text.rstrip('.')
    words = text.split()
    name = ''
    for word in words:
        if len(name) + len(word) + 1 > 52:
            break
        name = f'{name} {word}'.strip()
    return name.rstrip('.,:;') + '…'


def import_key(full_name, about):
    raw = f"{_norm(full_name)}|{_norm(about)[:160]}"
    return hashlib.sha1(raw.encode('utf-8')).hexdigest()


# -------------------------------------------------------------- rasmlar

def _images_by_cell(ws):
    """{(qator, ustun): rasm} — har katakdagi birinchi rasm."""
    found = {}
    for image in getattr(ws, '_images', []):
        anchor = getattr(image.anchor, '_from', None)
        if anchor is None:
            continue
        key = (anchor.row + 1, anchor.col + 1)
        found.setdefault(key, image)
    return found


def _image_file(image, stem):
    try:
        data = image._data()
    except Exception:  # noqa: BLE001 — buzilgan rasm butun importni to'xtatmasin
        return None
    if not data:
        return None
    extension = (getattr(image, 'format', '') or 'png').lower().replace('jpeg', 'jpg')
    return ContentFile(data, name=f'{stem}.{extension}')


# -------------------------------------------------------------- import

def import_workbook(file):
    """Faylni o'qiydi va bazaga yozadi. Natija: qancha qo'shildi, takror, xatolar."""
    try:
        workbook = load_workbook(file, data_only=True)
    except Exception:  # noqa: BLE001
        return {'detail': "Fayl o'qilmadi. .xlsx formatidagi Excel fayl yuboring."}

    ws = workbook.worksheets[0]
    header, columns = _find_header(ws)
    if not header:
        return {'detail': "«F.I.Sh» ustuni topilmadi. Birinchi qatorlarda ustun nomlari bo'lishi kerak."}
    if ws.max_row - header > MAX_ROWS:
        return {'detail': f"Juda ko'p qator — bir martada {MAX_ROWS} tagacha."}

    images = _images_by_cell(ws)
    existing = set(OfficeStartup.objects.values_list('import_key', flat=True))
    created = duplicates = empty = 0
    problems = []

    def value(row, kind):
        column = columns.get(kind)
        return ws.cell(row, column).value if column else None

    for row in range(header + 1, ws.max_row + 1):
        full_name = _clean(value(row, 'full_name'))
        about = str(value(row, 'about') or '').strip()
        if not full_name:
            empty += 1
            continue

        key = import_key(full_name, about)
        if key in existing:
            duplicates += 1
            continue

        birth_date, age = parse_age(value(row, 'age'))
        name = _clean(value(row, 'name')) or project_name(about) or f"{full_name.split()[0]} loyihasi"

        item = OfficeStartup(
            full_name=full_name[:150],
            district=parse_district(value(row, 'district')),
            birth_date=birth_date,
            age=age,
            phone=parse_phone(value(row, 'phone')),
            telegram=parse_telegram(value(row, 'telegram')),
            name=name[:200],
            about=about,
            sphere=parse_sphere(value(row, 'sphere')),
            stage=parse_stage(value(row, 'stage')),
            import_key=key,
        )

        try:
            with transaction.atomic():
                photo = images.get((row, columns.get('photo')))
                if photo is not None:
                    file = _image_file(photo, f'asoschi-{row}')
                    if file:
                        # Faylni maydonga beramiz — saqlashda avval siqiladi, keyin yoziladi
                        item.photo = file
                picture = images.get((row, columns.get('project_image')))
                if picture is not None:
                    file = _image_file(picture, f'loyiha-{row}')
                    if file:
                        item.project_image = file
                item.save()
        except Exception as error:  # noqa: BLE001 — bitta qator butun faylni to'xtatmasin
            problems.append(f"{row}-qator ({full_name}): {error}")
            continue

        existing.add(key)
        created += 1

    return {
        'created': created,
        'duplicates': duplicates,
        'empty': empty,
        'problems': problems[:30],
        'columns': sorted(get_column_letter(c) + ':' + k for k, c in columns.items()),
    }
