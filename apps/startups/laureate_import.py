"""«Mirzo Ulug'bek vorislari» g'oliblarini Excel'dan yuklash.

O'qiladigan ustunlar (sarlavhasi bo'yicha topiladi, tartibi muhim emas):

    Yil, F.I.Sh, Tuman, Loyiha nomi, Loyiha haqida, Tug'ilgan sana (yoki
    «Kun, oy, yil»), Yoshi, Ta'lim muassasasi, Kurs, Respublika g'olibi, Rasmi

Qolganlari (telefon, fudbolka o'lchami va h.k.) e'tiborsiz qoladi. «Yil»
ustuni bo'lmasa — yil sarlavhadagi matndan («MA'LUMOT - 2025») olinadi.
Aynan o'sha odamning o'sha yilgi o'sha loyihasi qayta yuklansa — o'tkazib yuboriladi.
"""

import hashlib
import re

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.core.constants import district_label

from .models import Laureate
from .office_import import _clean, _image_file, _images_by_cell, _norm, parse_age

MAX_ROWS = 3000
YEAR = re.compile(r'\b(20[2-4]\d)\b')


def _column_kind(header):
    h = _norm(header)
    if not h:
        return None
    if h in ('yil', 'yili', 'tanlov yili') or h.startswith('yil'):
        return 'year'
    if 'rasm' in h or 'foto' in h or 'surat' in h:
        return 'photo'
    if 'loyiha nomi' in h or h == 'loyiha':
        return 'project_name'
    if 'loyiha haqida' in h or 'tavsif' in h:
        return 'about'
    if 'f.i.sh' in h or 'fish' in h or 'f.i.o' in h or h in ('fio', 'ism', 'ismi'):
        return 'full_name'
    if 'tuman' in h or 'hudud' in h:
        return 'district'
    if 'tugilgan' in h or 'kun, oy' in h or 'kun oy' in h:
        return 'birth'
    if 'yosh' in h:
        return 'age'
    if 'talim' in h or 'oqish joyi' in h or 'muassasa' in h:
        return 'education'
    if 'kurs' in h or 'sinf' in h:
        return 'course'
    if 'respublika' in h:
        return 'national'
    return None


def _find_header(ws):
    for row in range(1, min(ws.max_row, 8) + 1):
        columns = {}
        for cell in ws[row]:
            kind = _column_kind(cell.value)
            if kind and kind not in columns:
                columns[kind] = cell.column
        if 'full_name' in columns and 'project_name' in columns:
            return row, columns
    return None, {}


def _sheet_year(ws, header):
    """Ustun bo'lmasa: yuqoridagi sarlavha matnidagi yil («MA'LUMOT - 2025»)."""
    for row in range(1, header):
        for cell in ws[row]:
            match = YEAR.search(str(cell.value or ''))
            if match:
                return int(match.group(1))
    return timezone.localdate().year


def parse_year(raw, fallback):
    match = YEAR.search(str(raw or ''))
    return int(match.group(1)) if match else fallback


def parse_district(raw):
    text = _clean(raw)
    if not text:
        return ''
    text = re.sub(r'\s+tuman$', ' tumani', text).replace('Samarqand shahar', 'Samarqand shahri')
    return district_label(text) or district_label(f'{text} tumani') or text[:100]


def parse_course(raw):
    value = _clean(raw).replace('синф', 'sinf')
    if isinstance(raw, (int, float)) or value.isdigit():
        return f"{int(float(value))}-kurs"
    return value[:40]


def parse_yes(raw):
    return _norm(raw) in ('ha', 'xa', 'yes', '1', 'true', '+', 'respublika golibi')


def import_key(full_name, project_name, year):
    raw = f"{_norm(full_name)}|{_norm(project_name)}|{year}"
    return hashlib.sha1(raw.encode('utf-8')).hexdigest()


def import_workbook(file):
    try:
        workbook = load_workbook(file, data_only=True)
    except Exception:  # noqa: BLE001
        return {'detail': "Fayl o'qilmadi. .xlsx formatidagi Excel fayl yuboring."}

    ws = workbook.worksheets[0]
    header, columns = _find_header(ws)
    if not header:
        return {'detail': "«F.I.Sh» va «Loyiha nomi» ustunlari topilmadi."}
    if ws.max_row - header > MAX_ROWS:
        return {'detail': f"Juda ko'p qator — bir martada {MAX_ROWS} tagacha."}

    default_year = _sheet_year(ws, header)
    images = _images_by_cell(ws)
    existing = set(Laureate.objects.values_list('import_key', flat=True))
    created = duplicates = empty = 0
    problems = []

    def value(row, kind):
        column = columns.get(kind)
        return ws.cell(row, column).value if column else None

    for row in range(header + 1, ws.max_row + 1):
        full_name = _clean(value(row, 'full_name'))
        project_name = _clean(value(row, 'project_name'))
        # «Izoh:» kabi pastdagi izoh qatorlari va bo'sh qatorlar
        if not full_name or not project_name:
            empty += 1
            continue

        year = parse_year(value(row, 'year'), default_year)
        key = import_key(full_name, project_name, year)
        if key in existing:
            duplicates += 1
            continue

        birth_date, birth_age = parse_age(value(row, 'birth'))
        _date, age = parse_age(value(row, 'age'))
        item = Laureate(
            full_name=full_name[:150],
            year=year,
            district=parse_district(value(row, 'district')),
            birth_date=birth_date,
            age=age or birth_age,
            education_place=_clean(value(row, 'education'))[:250],
            course=parse_course(value(row, 'course')),
            project_name=project_name[:250],
            about=str(value(row, 'about') or '').strip(),
            is_national_winner=parse_yes(value(row, 'national')),
            import_key=key,
        )

        try:
            with transaction.atomic():
                photo = images.get((row, columns.get('photo')))
                if photo is not None:
                    file = _image_file(photo, f'voris-{year}-{row}')
                    if file:
                        item.photo = file
                item.save()
        except Exception as error:  # noqa: BLE001 — bitta qator butun faylni to'xtatmasin
            problems.append(f"{row}-qator ({full_name}): {error}")
            continue

        existing.add(key)
        created += 1

    return {
        'created': created,
        'duplicates': duplicates,
        'empty': empty,
        'problems': problems[:30],
        'columns': sorted(get_column_letter(c) + ':' + k for k, c in columns.items()),
    }
