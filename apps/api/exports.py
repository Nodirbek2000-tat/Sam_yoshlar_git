"""Panel: bo'limlarni Excel (.xlsx) faylga yuklab olish.

Har bir bo'lim — alohida fayl, ichida bir nechta varaq: masalan
«Tashabbuslar» faylida tashabbuslarning o'zi, kim ovoz bergani va kim
taklif yozgani. Davr (`dan`/`gacha`) tanlansa, har bir varaq o'z sanasi
bo'yicha filtrlanadi — «shu oyda kim nima qildi» degan savolga javob.

Fayl `write_only` rejimida quriladi: qatorlar xotirada to'planmaydi,
o'n minglab yozuv bo'lsa ham server qiynalmaydi.
"""

from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth import get_user_model
from django.db.models import Count
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.abroad.models import Peer
from apps.accounts.models import Role
from apps.business.models import BusinessProfile, Product
from apps.content.models import Event, EventRegistration
from apps.initiatives.models import (Initiative, InitiativeComment, InitiativeVote,
                                     Organization, Problem, Solution)
from apps.startups.models import Startup

User = get_user_model()

#: Excel katakchasiga sig'adigan eng uzun matn
CELL_LIMIT = 32000

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1F7A4D')
HEADER_ALIGN = Alignment(vertical='center', wrap_text=True)


class Period:
    """Tanlangan davr: `dan` va `gacha` (ikkalasi ham kun, ikkalasi ham ixtiyoriy)."""

    def __init__(self, start=None, end=None):
        self.start, self.end = start, end

    @classmethod
    def from_query(cls, params):
        """`?dan=2026-09-01&gacha=2026-09-30` — noto'g'ri bo'lsa `ValueError`."""
        values = []
        for name in ('dan', 'gacha'):
            raw = (params.get(name) or '').strip()
            value = parse_date(raw) if raw else None
            if raw and value is None:
                raise ValueError("Sana noto'g'ri kiritilgan.")
            values.append(value)
        if values[0] and values[1] and values[0] > values[1]:
            raise ValueError("«Dan» sanasi «gacha» sanasidan keyin bo'lmasligi kerak.")
        return cls(*values)

    def apply(self, queryset, field='created_at'):
        tz = timezone.get_current_timezone()
        if self.start:
            queryset = queryset.filter(**{
                f'{field}__gte': datetime.combine(self.start, time.min, tzinfo=tz)})
        if self.end:
            queryset = queryset.filter(**{
                f'{field}__lte': datetime.combine(self.end, time.max, tzinfo=tz)})
        return queryset

    @property
    def suffix(self):
        """Fayl nomi uchun: `2026-09-01_2026-09-30` yoki bugungi sana."""
        if self.start or self.end:
            return f"{self.start or 'boshidan'}_{self.end or 'bugungacha'}"
        return str(timezone.localdate())


# --- Qiymatlarni Excel'ga moslash ---------------------------------------------

def _cell(sheet, value):
    """Python qiymatini Excel katakchasiga aylantiradi."""
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return "Ha" if value else "Yo'q"
    if isinstance(value, datetime):
        cell = WriteOnlyCell(sheet, timezone.localtime(value).replace(tzinfo=None, microsecond=0))
        cell.number_format = 'DD.MM.YYYY HH:MM'
        return cell
    if isinstance(value, date):
        cell = WriteOnlyCell(sheet, value)
        cell.number_format = 'DD.MM.YYYY'
        return cell
    if isinstance(value, (int, float, Decimal)):
        return value

    text = ILLEGAL_CHARACTERS_RE.sub('', str(value))[:CELL_LIMIT]
    if text.startswith('='):
        # Foydalanuvchi yozgan «=...» formula bo'lib ishlab ketmasin — oddiy matn
        cell = WriteOnlyCell(sheet, text)
        cell.data_type = 's'
        return cell
    return text


def _person_email(email):
    """Telegram orqali kirganlarning ichki manzili (`tg…@telegram.local`) ko'rsatilmaydi."""
    return '' if not email or email.endswith('@telegram.local') else email


def _link(path):
    return f"{settings.SITE_URL}{path}" if settings.SITE_URL else path


def _file_url(field):
    return _link(field.url) if field else ''


def _counts(queryset, field):
    """`{foydalanuvchi_id: soni}` — bitta GROUP BY so'rovida."""
    return Counter(dict(
        queryset.exclude(**{f'{field}__isnull': True}).order_by()
        .values_list(field).annotate(total=Count('id'))))


# --- Varaqlar -----------------------------------------------------------------
#
# Har bir varaq: (nomi, [(ustun, kengligi), ...], qatorlar generatori)

def initiatives_sheets(period):
    initiatives = (period.apply(Initiative.objects.select_related('author'))
                   .annotate(comment_total=Count('comments', distinct=True))
                   .order_by('-vote_count', '-created_at'))

    def rows():
        for number, item in enumerate(initiatives.iterator(chunk_size=1000), 1):
            author = item.author
            yield (number, item.title, item.get_direction_display(), item.get_kind_display(),
                   item.author_name, item.author_phone or (author.phone if author else ''),
                   _person_email(item.author_email), author.district if author else '',
                   item.get_region_display(), item.vote_count, item.comment_total,
                   item.get_status_display(), item.is_published, item.summary,
                   item.description, item.expected_result, item.admin_note, item.created_at,
                   _link(f'/tashabbuslar/{item.pk}'))

    votes = (period.apply(InitiativeVote.objects.select_related('initiative', 'user'))
             .order_by('-created_at'))

    def vote_rows():
        for number, vote in enumerate(votes.iterator(chunk_size=2000), 1):
            user = vote.user
            yield (number, vote.initiative.title,
                   user.full_name if user else "Mehmon (ro'yxatdan o'tmagan)",
                   user.phone if user else '', user.district if user else '', vote.created_at)

    comments = (period.apply(InitiativeComment.objects.select_related('initiative', 'author'))
                .order_by('-created_at'))

    def comment_rows():
        for number, comment in enumerate(comments.iterator(chunk_size=2000), 1):
            author = comment.author
            yield (number, comment.initiative.title, comment.author_name,
                   author.phone if author else '', author.district if author else '',
                   comment.text, comment.is_published, comment.created_at)

    return [
        ("Tashabbuslar", [
            ("№", 6), ("Sarlavha", 40), ("Yo'nalish", 22), ("Turi", 14), ("Muallif", 26),
            ("Telefon", 16), ("Email", 24), ("Tuman / shahar", 20), ("Viloyat", 18),
            ("Ovozlar", 9), ("Takliflar", 10), ("Holat", 16), ("Saytda ko'rinadi", 10),
            ("Qisqa mazmun", 40), ("Batafsil", 60), ("Kutilayotgan natija", 40),
            ("Admin izohi", 30), ("Qo'shilgan", 17), ("Havola", 40),
        ], rows()),
        ("Ovoz berganlar", [
            ("№", 6), ("Tashabbus", 40), ("Kim", 28), ("Telefon", 16), ("Tuman / shahar", 20),
            ("Sana", 17),
        ], vote_rows()),
        ("Takliflar", [
            ("№", 6), ("Tashabbus", 40), ("Kim yozdi", 26), ("Telefon", 16),
            ("Tuman / shahar", 20), ("Taklif matni", 70), ("Saytda ko'rinadi", 10), ("Sana", 17),
        ], comment_rows()),
    ]


def problems_sheets(period):
    problems = (period.apply(Problem.objects.select_related('organization'))
                .annotate(solution_total=Count('solutions', distinct=True))
                .order_by('-created_at'))

    def rows():
        for number, problem in enumerate(problems.iterator(chunk_size=1000), 1):
            org = problem.organization
            yield (number, org.name, org.get_sphere_display(), problem.get_category_display(),
                   problem.description, problem.solution_total, problem.get_status_display(),
                   problem.is_published, org.contact_person, org.phone, org.email,
                   problem.created_at, _link(f'/tashabbuslar/muammolar/{problem.pk}'))

    solutions = (period.apply(Solution.objects.select_related('problem__organization', 'author'))
                 .order_by('-created_at'))

    def solution_rows():
        for number, item in enumerate(solutions.iterator(chunk_size=1000), 1):
            problem, author = item.problem, item.author
            yield (number, problem.organization.name, problem.get_category_display(),
                   problem.description[:200], item.author_name,
                   item.author_phone or (author.phone if author else ''),
                   _person_email(item.author_email), author.district if author else '',
                   item.title, item.description, item.technologies, item.expected_result,
                   _file_url(item.attachment), item.like_count, item.get_status_display(),
                   item.admin_note, item.created_at)

    # Tashkilotlar — ma'lumotnoma: davrga qaramay hammasi
    organizations = (Organization.objects.select_related('user')
                     .annotate(problem_total=Count('problems', distinct=True)).order_by('name'))

    def organization_rows():
        for number, org in enumerate(organizations.iterator(chunk_size=1000), 1):
            account = org.user
            yield (number, org.name, org.get_sphere_display(), org.contact_person, org.phone,
                   org.email, org.get_region_display(), org.employees, org.problem_total,
                   _person_email(account.email) if account else '',
                   bool(account and account.telegram_id), org.created_at)

    return [
        ("Muammolar", [
            ("№", 6), ("Tashkilot", 32), ("Soha", 20), ("Kategoriya", 20),
            ("Muammo tavsifi", 70), ("Takliflar soni", 10), ("Holat", 16),
            ("Saytda ko'rinadi", 10), ("Aloqa shaxsi", 24), ("Telefon", 16), ("Email", 24),
            ("Qo'shilgan", 17), ("Havola", 40),
        ], rows()),
        ("Yechim takliflari", [
            ("№", 6), ("Tashkilot", 30), ("Muammo kategoriyasi", 20), ("Muammo", 40),
            ("Kim taklif qildi", 26), ("Telefon", 16), ("Email", 24), ("Tuman / shahar", 20),
            ("Yechim nomi", 32), ("Yechim tavsifi", 60), ("Texnologiyalar", 28),
            ("Kutilayotgan natija", 40), ("Fayl", 36), ("Layklar", 9), ("Holat", 16),
            ("Admin izohi", 30), ("Sana", 17),
        ], solution_rows()),
        ("Tashkilotlar", [
            ("№", 6), ("Nomi", 32), ("Soha", 20), ("Aloqa shaxsi", 24), ("Telefon", 16),
            ("Email", 24), ("Viloyat", 18), ("Xodimlar", 10), ("Muammolar soni", 10),
            ("Kirish hisobi", 28), ("Telegram ulangan", 10), ("Qo'shilgan", 17),
        ], organization_rows()),
    ]


def businesses_sheets(period):
    businesses = (period.apply(BusinessProfile.objects.select_related('user'))
                  .annotate(product_total=Count('products', distinct=True))
                  .order_by('-created_at'))

    def rows():
        for number, item in enumerate(businesses.iterator(chunk_size=1000), 1):
            owner = item.user
            yield (number, item.name, item.get_sphere_display(), owner.full_name,
                   owner.phone, item.phone, _person_email(item.email), item.stir,
                   item.founded_year, item.employees, item.get_region_display(), item.district,
                   item.address, item.telegram, item.instagram, item.website,
                   item.product_total, item.get_status_display(), item.is_public,
                   item.description, item.created_at, _link(f'/tadbirkorlar/{item.pk}'))

    products = (period.apply(Product.objects.select_related('business'))
                .order_by('business__name', 'name'))

    def product_rows():
        for number, item in enumerate(products.iterator(chunk_size=2000), 1):
            yield (number, item.business.name, item.name, item.price, item.unit,
                   item.is_active, item.description, item.created_at)

    return [
        ("Tadbirkorlar", [
            ("№", 6), ("Korxona", 32), ("Soha", 22), ("Egasi", 26), ("Egasining telefoni", 16),
            ("Korxona telefoni", 16), ("Email", 24), ("STIR", 12), ("Tashkil topgan yili", 10),
            ("Xodimlar", 10), ("Viloyat", 18), ("Tuman / shahar", 20), ("Manzil", 30),
            ("Telegram", 18), ("Instagram", 18), ("Veb-sayt", 26), ("Mahsulotlar", 11),
            ("Holat", 16), ("Saytda ko'rinadi", 10), ("Biznes haqida", 60), ("Qo'shilgan", 17),
            ("Havola", 40),
        ], rows()),
        ("Mahsulotlar", [
            ("№", 6), ("Korxona", 32), ("Mahsulot", 32), ("Narxi (so'm)", 14),
            ("O'lchov birligi", 12), ("Faol", 8), ("Tavsif", 60), ("Qo'shilgan", 17),
        ], product_rows()),
    ]


def startups_sheets(period):
    startups = period.apply(Startup.objects.all()).order_by('-created_at')

    def rows():
        for number, item in enumerate(startups.iterator(chunk_size=1000), 1):
            yield (number, item.name, item.get_sphere_display(), item.get_stage_display(),
                   item.full_name, item.phone, _person_email(item.email), item.birth_date,
                   item.get_region_display(), item.district, item.team_size,
                   item.needed_investment, item.about, item.problem_solved, item.website,
                   _file_url(item.pitch_file), item.get_status_display(), item.is_public,
                   item.admin_note, item.created_at, _link(f'/startaplar/{item.pk}'))

    return [
        ("Startaplar", [
            ("№", 6), ("Startap", 30), ("Yo'nalish", 20), ("Bosqich", 16), ("Muallif", 26),
            ("Telefon", 16), ("Email", 24), ("Tug'ilgan sana", 12), ("Viloyat", 18),
            ("Tuman / shahar", 20), ("Jamoa", 8), ("Kerakli investitsiya (so'm)", 16),
            ("Startap haqida", 60), ("Qanday muammoni hal qiladi", 50), ("Veb-sayt", 26),
            ("Pitch fayl", 36), ("Holat", 16), ("Saytda ko'rinadi", 10), ("Admin izohi", 30),
            ("Qo'shilgan", 17), ("Havola", 40),
        ], rows()),
    ]


def peers_sheets(period):
    peers = period.apply(Peer.objects.all()).order_by('country', 'full_name')

    def rows():
        for number, item in enumerate(peers.iterator(chunk_size=1000), 1):
            yield (number, item.full_name, item.get_country_display(), item.city,
                   item.get_home_region_display(), item.get_purpose_display(), item.institution,
                   item.field, item.since_year, item.achievements, item.about, item.can_help,
                   item.telegram, item.instagram, _person_email(item.email), item.phone,
                   item.get_status_display(), item.is_published, item.admin_note,
                   item.created_at, _link(f'/tengdoshlar/{item.pk}'))

    return [
        ("Tengdoshlar", [
            ("№", 6), ("F.I.O.", 28), ("Davlat", 18), ("Shahar", 16),
            ("O'zbekistondagi hududi", 20), ("Nima qilyapti", 18),
            ("Universitet / kompaniya", 30), ("Yo'nalish / kasb", 24), ("Qaysi yildan", 10),
            ("Yutuqlari", 40), ("O'zi haqida", 50), ("Nimada yordam bera oladi", 50),
            ("Telegram", 18), ("Instagram", 18), ("Email", 24), ("Telefon", 16),
            ("Holat", 16), ("Saytda ko'rinadi", 10), ("Admin izohi", 30), ("Qo'shilgan", 17),
            ("Havola", 40),
        ], rows()),
    ]


def users_sheets(period):
    """Foydalanuvchilar va ularning faolligi — «kim nima qildi».

    Faollik sonlari tanlangan davr ichida hisoblanadi.
    """
    initiatives = _counts(period.apply(Initiative.objects), 'author')
    votes = _counts(period.apply(InitiativeVote.objects), 'user')
    comments = _counts(period.apply(InitiativeComment.objects), 'author')
    solutions = _counts(period.apply(Solution.objects), 'author')
    startups = _counts(period.apply(Startup.objects), 'user')
    events = _counts(period.apply(EventRegistration.objects.filter(is_cancelled=False)), 'user')
    businesses = dict(BusinessProfile.objects.values_list('user', 'status'))
    business_status = dict(BusinessProfile._meta.get_field('status').choices)

    # Tashkilot va bosh adminlar — foydalanuvchi emas (panel ro'yxati bilan bir xil)
    users = (period.apply(User.objects.exclude(role=Role.ORGANIZATION).filter(is_superuser=False),
                          'date_joined')
             .order_by('-date_joined'))

    def rows():
        for number, user in enumerate(users.iterator(chunk_size=2000), 1):
            pk = user.pk
            yield (number, user.full_name, user.phone, _person_email(user.email),
                   user.get_role_display(), user.age, user.district, user.get_region_display(),
                   user.get_study_location_display(),
                   f"@{user.telegram_username}" if user.telegram_username else '',
                   user.date_joined, user.last_login, initiatives[pk], votes[pk], comments[pk],
                   solutions[pk], startups[pk], business_status.get(businesses.get(pk), ''),
                   events[pk])

    return [
        ("Foydalanuvchilar", [
            ("№", 6), ("F.I.O.", 28), ("Telefon", 16), ("Email", 24), ("Rol", 14), ("Yoshi", 7),
            ("Tuman / shahar", 20), ("Viloyat", 18), ("Qayerda ta'lim oladi", 16),
            ("Telegram", 18), ("Ro'yxatdan o'tgan", 17), ("Oxirgi kirish", 17),
            ("Tashabbuslari", 11), ("Bergan ovozlari", 11), ("Yozgan takliflari", 11),
            ("Muammoga yechimlari", 11), ("Startaplari", 11), ("Tadbirkor anketasi", 16),
            ("Tadbirlarga yozilgan", 11),
        ], rows()),
    ]


def events_sheets(period):
    selected = period.apply(Event.objects.all(), 'starts_at')
    events = (selected.annotate(registered=Count('registrations', distinct=True))
              .order_by('-starts_at'))

    def rows():
        for number, event in enumerate(events.iterator(chunk_size=1000), 1):
            yield (number, event.title, event.starts_at, event.ends_at, event.location,
                   event.get_region_display(), event.capacity, event.registered,
                   event.is_published, _link(f'/tadbirlar/{event.slug}'))

    registrations = (EventRegistration.objects.filter(event__in=selected.values('pk'))
                     .select_related('event', 'user').order_by('-event__starts_at', 'created_at'))

    def registration_rows():
        for number, item in enumerate(registrations.iterator(chunk_size=2000), 1):
            user = item.user
            yield (number, item.event.title, item.event.starts_at, user.full_name, user.phone,
                   user.district, user.age, item.created_at, item.is_cancelled, item.attended,
                   item.note)

    return [
        ("Tadbirlar", [
            ("№", 6), ("Tadbir", 40), ("Boshlanishi", 17), ("Tugashi", 17), ("Manzil", 30),
            ("Viloyat", 18), ("Limit", 8), ("Yozilganlar", 11), ("Saytda ko'rinadi", 10),
            ("Havola", 40),
        ], rows()),
        ("Ishtirokchilar", [
            ("№", 6), ("Tadbir", 40), ("Tadbir sanasi", 17), ("F.I.O.", 28), ("Telefon", 16),
            ("Tuman / shahar", 20), ("Yoshi", 7), ("Yozilgan", 17), ("Bekor qilgan", 10),
            ("Qatnashgan", 10), ("Izoh", 40),
        ], registration_rows()),
    ]


# --- Bo'limlar ro'yxati -------------------------------------------------------
#
# Kalit fayl nomiga va manzilga yoziladi (`/panel/eksport/tashabbuslar/`).
# `parts` — fayldagi varaqlar (panelda ko'rsatiladi), `count` — «nechta yozuv»
# uchun asosiy queryset va uning sana maydoni.

EXPORTS = {
    'tashabbuslar': {
        'title': "Yoshlar tashabbuslari",
        'description': "Tashabbuslar, kim ovoz bergani va kim taklif yozgani",
        'icon': 'spark',
        'sheets': initiatives_sheets,
        'parts': ["Tashabbuslar", "Ovoz berganlar", "Takliflar"],
        'count': lambda: (Initiative.objects, 'created_at'),
    },
    'muammolar': {
        'title': "Tashkilot muammolari",
        'description': "Muammolar, yoshlarning yechim takliflari va tashkilotlar ro'yxati",
        'icon': 'clipboard',
        'sheets': problems_sheets,
        'parts': ["Muammolar", "Yechim takliflari", "Tashkilotlar"],
        'count': lambda: (Problem.objects, 'created_at'),
    },
    'tadbirkorlar': {
        'title': "Tadbirkorlar",
        'description': "Korxonalar, egalari, aloqa ma'lumotlari va mahsulotlari",
        'icon': 'briefcase',
        'sheets': businesses_sheets,
        'parts': ["Tadbirkorlar", "Mahsulotlar"],
        'count': lambda: (BusinessProfile.objects, 'created_at'),
    },
    'startaplar': {
        'title': "Startaplar",
        'description': "Startaplar, mualliflari, bosqichi va kerakli investitsiya",
        'icon': 'rocket',
        'sheets': startups_sheets,
        'parts': ["Startaplar"],
        'count': lambda: (Startup.objects, 'created_at'),
    },
    'tengdoshlar': {
        'title': "Chet eldagi tengdoshlar",
        'description': "Qaysi davlatda, nima qilyapti, qanday yordam bera oladi",
        'icon': 'globe',
        'sheets': peers_sheets,
        'parts': ["Tengdoshlar"],
        'count': lambda: (Peer.objects, 'created_at'),
    },
    'foydalanuvchilar': {
        'title': "Foydalanuvchilar",
        'description': "Har bir foydalanuvchi va uning faolligi: tashabbus, ovoz, taklif, tadbir",
        'icon': 'users',
        'sheets': users_sheets,
        'parts': ["Foydalanuvchilar"],
        'count': lambda: (User.objects.exclude(role=Role.ORGANIZATION)
                          .filter(is_superuser=False), 'date_joined'),
    },
    'tadbirlar': {
        'title': "Tadbirlar",
        'description': "Tadbirlar va ularga yozilgan ishtirokchilar",
        'icon': 'calendar',
        'sheets': events_sheets,
        'parts': ["Tadbirlar", "Ishtirokchilar"],
        'count': lambda: (Event.objects, 'starts_at'),
    },
}


def export_count(key, period):
    queryset, field = EXPORTS[key]['count']()
    return period.apply(queryset, field).count()


def build_workbook(key, period):
    """Bo'limning Excel faylini bayt ko'rinishida qaytaradi."""
    book = Workbook(write_only=True)

    for title, columns, rows in EXPORTS[key]['sheets'](period):
        sheet = book.create_sheet(title)
        sheet.freeze_panes = 'A2'
        for index, (_, width) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        header = []
        for name, _ in columns:
            cell = WriteOnlyCell(sheet, name)
            cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGN
            header.append(cell)
        sheet.append(header)

        total = 0
        for row in rows:
            sheet.append([_cell(sheet, value) for value in row])
            total += 1
        # Sarlavhadagi filtr tugmalari — Excel'da darhol saralash mumkin
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{total + 1}"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
