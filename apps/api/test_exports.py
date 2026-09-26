"""Panel: bo'limlarni Excel'ga yuklab olish."""

from datetime import datetime, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.abroad.models import Peer
from apps.business.models import BusinessProfile, Product
from apps.content.models import Event, EventRegistration
from apps.initiatives.models import (Initiative, InitiativeComment, InitiativeVote,
                                     Organization, Problem, Solution)
from apps.startups.models import Startup

from .auth_views import tokens_for
from .exports import EXPORTS

User = get_user_model()
URL = '/api/v1/panel/eksport/'


def rows(book, sheet):
    return [list(row) for row in book[sheet].iter_rows(values_only=True)]


class ExportTests(TestCase):
    def setUp(self):
        admin = User.objects.create_superuser(email='a@test.uz', password='x', full_name="Admin")
        self.auth = {'HTTP_AUTHORIZATION': f"Bearer {tokens_for(admin)['access']}"}

        # Telegram orqali kirgan yosh — ichki email Excel'da ko'rinmasligi kerak
        self.yosh = User.objects.create_user(
            email='tg77@telegram.local', full_name="Ali Valiyev", phone='+998901234567',
            role='startupper', district="Urgut tumani", age=21)
        guest_voter = User.objects.create_user(email='b@test.uz', full_name="Bobur")

        self.idea = Initiative.objects.create(
            title="Yashil hovli", description="Daraxt ekamiz", author=self.yosh,
            author_name="Ali Valiyev", vote_count=2)
        InitiativeVote.objects.create(initiative=self.idea, user=guest_voter)
        InitiativeVote.objects.create(initiative=self.idea, session_key='mehmon-sessiya')
        InitiativeComment.objects.create(initiative=self.idea, author=self.yosh,
                                         author_name="Ali Valiyev", text="Men ham qatnashaman")

        account = User.objects.create_user(email='org@test.uz', full_name="Org",
                                           role='organization', telegram_id=555)
        org = Organization.objects.create(user=account, name="AgroTech", sphere='it',
                                          contact_person="Karim", phone='+998907654321')
        problem = Problem.objects.create(organization=org, category='main',
                                         description="Suv yetishmaydi")
        Solution.objects.create(problem=problem, author=self.yosh, author_name="Ali Valiyev",
                                title="Tomchilatib sug'orish", description="Arzon tizim")

        owner = User.objects.create_user(email='biz@test.uz', full_name="Dilshod",
                                         phone='+998911112233', role='entrepreneur')
        business = BusinessProfile.objects.create(user=owner, name="Non uyi", sphere='it')
        Product.objects.create(business=business, name="Patir", price=12000)

        Startup.objects.create(user=self.yosh, full_name="Ali Valiyev", phone='+998901234567',
                               email='ali@test.uz', region='samarqand', name="EcoApp",
                               sphere='it', about="Ekologik ilova")
        Peer.objects.create(full_name="Madina", country='usa', institution="MIT")

        event = Event.objects.create(title="Hakaton", description="…", location="Samarqand",
                                     starts_at=timezone.now() + timedelta(days=3))
        EventRegistration.objects.create(event=event, user=self.yosh)

    def download(self, key, query=''):
        response = self.client.get(f'{URL}{key}/{query}', **self.auth)
        self.assertEqual(response.status_code, 200, response.content[:200])
        return response, load_workbook(BytesIO(response.content))

    def test_list_shows_every_section_with_counts(self):
        data = self.client.get(URL, **self.auth).json()['results']
        counts = {row['key']: row['count'] for row in data}

        self.assertEqual(set(counts), set(EXPORTS))
        self.assertEqual(counts, {
            'tashabbuslar': 1, 'muammolar': 1, 'tadbirkorlar': 1, 'startaplar': 1,
            'tengdoshlar': 1, 'tadbirlar': 1,
            # Tashkilot va bosh admin foydalanuvchi sifatida sanalmaydi
            'foydalanuvchilar': 3,
        })

    def test_every_section_downloads_as_excel(self):
        for key in EXPORTS:
            with self.subTest(key=key):
                response, book = self.download(key)
                self.assertTrue(response['Content-Type'].startswith(
                    'application/vnd.openxmlformats-officedocument.spreadsheetml'))
                self.assertIn(f'filename="{key}_', response['Content-Disposition'])
                self.assertEqual(book.sheetnames, EXPORTS[key]['parts'])
                for sheet in book.worksheets:
                    self.assertEqual(sheet['A1'].value, "№")
                    self.assertEqual(sheet.freeze_panes, 'A2')

    def test_initiatives_show_who_voted_and_commented(self):
        _, book = self.download('tashabbuslar')
        self.assertEqual(book.sheetnames, ["Tashabbuslar", "Ovoz berganlar", "Takliflar"])

        header, idea = rows(book, "Tashabbuslar")[:2]
        row = dict(zip(header, idea))
        self.assertEqual((row["Sarlavha"], row["Muallif"], row["Ovozlar"], row["Takliflar"]),
                         ("Yashil hovli", "Ali Valiyev", 2, 1))
        self.assertEqual(row["Telefon"], '+998901234567')      # anketada yo'q — hisobdan
        self.assertEqual(row["Tuman / shahar"], "Urgut tumani")
        self.assertIsNone(row["Email"])                        # tg…@telegram.local yashirin
        self.assertTrue(row["Havola"].endswith(f'/tashabbuslar/{self.idea.pk}'))
        self.assertIsInstance(row["Qo'shilgan"], datetime)

        voters = {line[2] for line in rows(book, "Ovoz berganlar")[1:]}
        self.assertEqual(voters, {"Bobur", "Mehmon (ro'yxatdan o'tmagan)"})
        self.assertEqual(rows(book, "Takliflar")[1][5], "Men ham qatnashaman")

    def test_problems_include_solutions_and_organizations(self):
        _, book = self.download('muammolar')
        self.assertEqual(book.sheetnames, ["Muammolar", "Yechim takliflari", "Tashkilotlar"])
        solution = dict(zip(*rows(book, "Yechim takliflari")[:2]))
        self.assertEqual((solution["Tashkilot"], solution["Kim taklif qildi"],
                          solution["Yechim nomi"]),
                         ("AgroTech", "Ali Valiyev", "Tomchilatib sug'orish"))
        org = dict(zip(*rows(book, "Tashkilotlar")[:2]))
        self.assertEqual((org["Nomi"], org["Muammolar soni"], org["Telegram ulangan"]),
                         ("AgroTech", 1, "Ha"))

    def test_users_sheet_shows_what_each_person_did(self):
        _, book = self.download('foydalanuvchilar')
        header, *people = rows(book, "Foydalanuvchilar")
        by_name = {row[1]: dict(zip(header, row)) for row in people}

        self.assertNotIn("Org", by_name)          # tashkilot
        self.assertNotIn("Admin", by_name)        # bosh admin
        ali = by_name["Ali Valiyev"]
        self.assertEqual(
            (ali["Tashabbuslari"], ali["Yozgan takliflari"], ali["Muammoga yechimlari"],
             ali["Startaplari"], ali["Tadbirlarga yozilgan"], ali["Bergan ovozlari"]),
            (1, 1, 1, 1, 1, 0))
        self.assertEqual(by_name["Bobur"]["Bergan ovozlari"], 1)
        self.assertEqual(by_name["Dilshod"]["Tadbirkor anketasi"], "Kutilmoqda")

    def test_period_filters_by_date(self):
        old = timezone.now() - timedelta(days=400)
        Initiative.objects.create(title="Eski g'oya", description="…", author_name="X")
        Initiative.objects.filter(title="Eski g'oya").update(created_at=old)

        today = timezone.localdate()
        query = f'?dan={today - timedelta(days=30)}&gacha={today}'
        counts = {row['key']: row['count'] for row in
                  self.client.get(URL + query, **self.auth).json()['results']}
        self.assertEqual(counts['tashabbuslar'], 1)

        response, book = self.download('tashabbuslar', query)
        titles = [row[1] for row in rows(book, "Tashabbuslar")[1:]]
        self.assertEqual(titles, ["Yashil hovli"])
        self.assertIn(str(today), response['Content-Disposition'])

    def test_user_text_never_becomes_a_formula(self):
        self.idea.title = '=HYPERLINK("http://yomon.uz","Bosing")'
        self.idea.save()
        _, book = self.download('tashabbuslar')
        cell = book["Tashabbuslar"]['B2']
        self.assertEqual(cell.data_type, 's')
        self.assertEqual(cell.value, '=HYPERLINK("http://yomon.uz","Bosing")')

    def test_bad_requests(self):
        self.assertEqual(self.client.get(f'{URL}?dan=kecha', **self.auth).status_code, 400)
        self.assertEqual(
            self.client.get(f'{URL}tashabbuslar/?dan=2026-09-30&gacha=2026-09-01',
                            **self.auth).status_code, 400)
        self.assertEqual(self.client.get(f'{URL}bunday-bolim-yoq/', **self.auth).status_code, 404)

    def test_only_panel_admin(self):
        auth = {'HTTP_AUTHORIZATION': f"Bearer {tokens_for(self.yosh)['access']}"}
        self.assertEqual(self.client.get(URL, **auth).status_code, 404)
        self.assertEqual(self.client.get(f'{URL}foydalanuvchilar/', **auth).status_code, 404)
        self.assertEqual(self.client.get(f'{URL}foydalanuvchilar/').status_code, 404)
